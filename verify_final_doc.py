import docx, os
from docx.oxml.ns import qn
import zipfile
from xml.etree import ElementTree as ET

doc_path = r'E:\VanDe_AI\docs\BAO_CAO_DO_AN_HUIT.docx'
print(f"File exists: {os.path.exists(doc_path)}, size: {os.path.getsize(doc_path):,} bytes")

doc = docx.Document(doc_path)
print(f"Paragraphs: {len(doc.paragraphs)}")
print(f"Tables: {len(doc.tables)}")
print(f"Sections: {len(doc.sections)}")

# Inspect sections
for i, s in enumerate(doc.sections):
    h_text = [p.text for p in s.header.paragraphs if p.text.strip()]
    f_text = [p.text for p in s.footer.paragraphs if p.text.strip()]
    print(f"=== Section {i} ===")
    print(f"  Margins: top={s.top_margin.pt}, bottom={s.bottom_margin.pt}, left={s.left_margin.pt}, right={s.right_margin.pt}")
    print(f"  Header text: {h_text}")
    print(f"  Footer text: {f_text}")

# Count headings
h1_list = [p.text for p in doc.paragraphs if p.style.name == 'Heading 1']
h2_list = [p.text for p in doc.paragraphs if p.style.name == 'Heading 2']
h3_list = [p.text for p in doc.paragraphs if p.style.name == 'Heading 3']
print(f"\nHeading 1 count: {len(h1_list)}")
for h in h1_list:
    print(f"  H1: {h}")
print(f"\nHeading 2 count: {len(h2_list)}")
for h in h2_list:
    print(f"  H2: {h}")
print(f"\nHeading 3 count: {len(h3_list)}")
for h in h3_list:
    print(f"  H3: {h}")

# Inspect OMML equations
z = zipfile.ZipFile(doc_path)
doc_xml = z.read('word/document.xml').decode('utf-8')
tree = ET.fromstring(doc_xml)
omath_elems = tree.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath')
print(f"\nTotal OMML equations found: {len(omath_elems)}")

# Inspect media
media_files = [f for f in z.namelist() if f.startswith('word/media/')]
print(f"Total media/images embedded: {len(media_files)}")
for m in media_files:
    info = z.getinfo(m)
    print(f"  Media: {m} ({info.file_size:,} bytes)")

# Inspect settings
if 'word/settings.xml' in z.namelist():
    settings_xml = z.read('word/settings.xml').decode('utf-8')
    print(f"updateFields in settings.xml: {'updateFields' in settings_xml}")
